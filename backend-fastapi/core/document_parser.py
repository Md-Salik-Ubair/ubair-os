import os
import re
from typing import List, Dict, Any, Optional

# Safe Library Imports
try:
    import pandas as pd
    _pandas_available = True
except ImportError:
    _pandas_available = False

try:
    from pypdf import PdfReader
    _pypdf_available = True
except ImportError:
    _pypdf_available = False

try:
    import docx
    _docx_available = True
except ImportError:
    _docx_available = False


class DocumentProcessor:
    """
    Enterprise Universal Ingestion & Smart Chunker Engine for Ubair OS:
    - Multi-Format Parser (PDF, Excel, CSV, DOCX, Code, Markdown, Structured Text)
    - Auto-Healing File Descriptors (Immune to Windows OS File-Lock WinError 32)
    - Multi-Encoding Auto-Fallback (UTF-8, UTF-16 LE, Latin-1)
    - Tabulate-Independent Resilient Table Serializer (Zero-Crash Guarantee)
    - Semantic Sentence-Boundary Chunker for High-Fidelity Vector Retrieval
    - Anti-DoS Chunk Capping & Null-Byte Sanitization
    """
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 150, max_chunks_per_file: int = 200):
        self.chunk_size = chunk_size
        self.overlap = chunk_overlap
        self.max_chunks_per_file = max_chunks_per_file

    def _sanitize_text(self, text: str) -> str:
        """Removes null bytes, zero-width spaces, and repetitive whitespace."""
        if not text:
            return ""
        cleaned = text.replace('\x00', ' ').replace('\u200b', '')
        cleaned = re.sub(r'[ \t]+', ' ', cleaned)
        cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
        return cleaned.strip()

    def process_file(
        self,
        file_path: str,
        original_filename: str,
        file_id: str,
        workspace_id: str,
        user_id: str = "anonymous"
    ) -> List[Dict[str, Any]]:
        """
        Universal Ingestion Pipeline:
        Parses document, sanitizes content, and emits isolated vector chunks.
        """
        if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
            return []

        ext = os.path.splitext(original_filename)[1].lower()
        extracted_text = ""

        try:
            if ext == '.pdf' and _pypdf_available:
                extracted_text = self._parse_pdf(file_path)
            elif ext in ['.xlsx', '.xls'] and _pandas_available:
                extracted_text = self._parse_excel(file_path)
            elif ext == '.csv' and _pandas_available:
                extracted_text = self._parse_csv(file_path)
            elif ext == '.docx' and _docx_available:
                extracted_text = self._parse_docx(file_path)
            else:
                extracted_text = self._parse_text(file_path)

            extracted_text = self._sanitize_text(extracted_text)

            if not extracted_text:
                if ext == '.pdf':
                    print(f"⚠️ [INGESTION NOTICE] '{original_filename}' appears to be a scanned/image PDF with no selectable text.")
                return []

            return self._chunk_text(
                text=extracted_text,
                filename=original_filename,
                file_id=file_id,
                workspace_id=workspace_id,
                user_id=user_id
            )

        except Exception as e:
            print(f"❌ [PARSER ERROR] Failed to parse '{original_filename}': {str(e)}")
            return []

    # ---------------------------------------------------------
    # SPECIFIC EXTRACTORS (WITH OS FILE-HANDLE SAFETY)
    # ---------------------------------------------------------
    def _parse_pdf(self, file_path: str) -> str:
        text_fragments = []
        try:
            # Explicit stream context guarantees zero file-lock on Windows
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for page_num, page in enumerate(reader.pages, start=1):
                    try:
                        page_text = page.extract_text() or ""
                        if page_text.strip():
                            text_fragments.append(f"[Page {page_num}]\n{page_text.strip()}")
                    except Exception:
                        continue
        except Exception as e:
            print(f"[PDF READ ERROR] {e}")
        return "\n\n".join(text_fragments)

    def _df_to_text(self, df: Any) -> str:
        """Converts DataFrame to text safely without relying on external tabulate."""
        try:
            return df.to_markdown(index=False)
        except Exception:
            pass
        try:
            return df.to_string(index=False)
        except Exception:
            return str(df.to_dict(orient="records"))

    def _parse_excel(self, file_path: str) -> str:
        sheets_data = []
        try:
            with open(file_path, "rb") as f:
                excel_file = pd.ExcelFile(f)
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet_name, nrows=500)
                    if not df.empty:
                        table_str = self._df_to_text(df)
                        sheets_data.append(f"### Sheet: {sheet_name} ###\n{table_str}")
        except Exception as e:
            print(f"[EXCEL READ ERROR] {e}")
        return "\n\n".join(sheets_data)

    def _parse_csv(self, file_path: str) -> str:
        try:
            blueprint = ""
            try:
                from core.analyst_profiler import AnalystProfiler
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f_raw:
                    profile = AnalystProfiler.profile_csv_content(f_raw.read(500000))
                    blueprint = AnalystProfiler.format_for_llm_context(profile, filename=os.path.basename(file_path)) + "\n\n"
            except Exception:
                blueprint = ""

            with open(file_path, "rb") as f:
                df = pd.read_csv(f, nrows=1000)
                return blueprint + self._df_to_text(df)
        except Exception:
            return self._parse_text(file_path)

    def _parse_docx(self, file_path: str) -> str:
        try:
            with open(file_path, "rb") as f:
                doc = docx.Document(f)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            paragraphs.append(row_text)
                            
                return "\n\n".join(paragraphs)
        except Exception as e:
            print(f"[DOCX READ ERROR] {e}")
            return ""

    def _parse_text(self, file_path: str) -> str:
        """Robust text file reader with multi-encoding fallback."""
        encodings = ['utf-8', 'utf-8-sig', 'utf-16', 'latin-1']
        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    return f.read(1_000_000)
            except (UnicodeDecodeError, UnicodeError):
                continue
            except Exception:
                break
        
        # Last resort byte salvage
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(1_000_000)
        except Exception:
            return ""

    # ---------------------------------------------------------
    # SMART SEMANTIC CHUNKER
    # ---------------------------------------------------------
    def _chunk_text(
        self,
        text: str,
        filename: str,
        file_id: str,
        workspace_id: str,
        user_id: str
    ) -> List[Dict[str, Any]]:
        chunks = []
        text_length = len(text)
        start = 0
        chunk_index = 1

        while start < text_length and chunk_index <= self.max_chunks_per_file:
            end = min(start + self.chunk_size, text_length)

            if end < text_length:
                # Priority 1: Double newline (Paragraph boundary)
                p_break = text.rfind('\n\n', start, end)
                if p_break != -1 and p_break > start + (self.chunk_size // 3):
                    end = p_break + 2
                else:
                    # Priority 2: Sentence boundary (. , ! , ?)
                    sentence_breaks = [
                        text.rfind('. ', start, end),
                        text.rfind('? ', start, end),
                        text.rfind('! ', start, end),
                        text.rfind('\n', start, end)
                    ]
                    valid_breaks = [b for b in sentence_breaks if b > start + (self.chunk_size // 2)]
                    if valid_breaks:
                        end = max(valid_breaks) + 1
                    else:
                        # Priority 3: Word space
                        last_space = text.rfind(' ', start, end)
                        if last_space != -1 and last_space > start + (self.chunk_size // 3):
                            end = last_space

            chunk_content = text[start:end].strip()

            if chunk_content and len(chunk_content) > 15:
                chunks.append({
                    "chunk_id": f"{file_id}_chunk_{chunk_index}",
                    "content": chunk_content,
                    "metadata": {
                        "user_id": str(user_id),
                        "workspace_id": str(workspace_id),
                        "file_id": file_id,
                        "filename": filename,
                        "chunk_index": chunk_index,
                        "char_count": len(chunk_content)
                    }
                })
                chunk_index += 1

            next_start = end - self.overlap
            if next_start <= start:
                next_start = end
            start = next_start

        return chunks


# Global Singleton Instance
document_processor = DocumentProcessor()